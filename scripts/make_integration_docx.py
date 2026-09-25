#!/usr/bin/env python
"""Generate the 课题一 technical integration Word document from the template.

Usage::

    python scripts/make_integration_docx.py \
        --template "【课题一】对接技术整理.docx" \
        --out "课题一_技术对接表_QACD.docx"

The script fills the four blank rows of the template's 技术集成 table with the
QACD technology points and appends a supplementary section (delivery status,
key numbers, resource envelope, mandatory platform safeguards, evidence
boundaries and open items).

It is deliberately reproducible: the document can be regenerated whenever the
numbers or the interface change. Keep `docs/05-课题一技术对接表.md` in sync.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

# Cell values for the four technology points, in template column order:
# 序号 | 所属子课题 | 技术点名称 | 功能简述 | 入参 | 出参 | 使用说明 |
# 资源需求 | 研究阶段 | 论文地址 | GitHub仓库/Huggingface | 对接负责人
SUBTopic = "子课题1"
REPO = "https://github.com/hlcccc/QACD"
PAPER = "待投稿（暂无公开链接）"
OWNER = "（待作者补充）"
STAGE = "论文撰写中"

ROWS = [
    # --- 技术点 1 -------------------------------------------------------
    [
        "QACD 回答错误风险评分器",
        "采用问题条件化原子主张分解 + 结构化证据校准双架构，覆盖 11 类主张类型与类型化验证路由，"
        "输出回答级错误风险分。黑盒后处理，不改动被评估模型。",
        '{"question": "string", "answer": "string", "image": "string(optional)", '
        '"threshold": "float(optional)"}',
        '{"risk_score": "float(0-1)", "is_high_risk": "bool", "num_claims": "int", '
        '"claims": "list[object]", "version": "string"}',
        "使用 GitHub 链接中提供的 qacd 包与 run.py，本地运行环境依赖见 GitHub requirements.txt",
        "单卡 40G 显存",
        STAGE,
        PAPER,
        REPO,
        OWNER,
    ],
    # --- 技术点 2 -------------------------------------------------------
    [
        "MVR 多视图重采样一致性通道",
        "采用同 prompt 多采样 + 图像文字机械核对双通道，覆盖 K 次采样的稳定性统计，"
        "输出 7 维一致性特征。需与主打分器联合使用。",
        '{"sampled_answers": "list[string]", "ocr_texts": "list[string]", "k": "int(optional)"}',
        '{"k_used": "int", "features": "object(7 维稳定性特征)"}',
        "使用 GitHub 链接中提供的 qacd.mvr_features 接口，K 建议取 3，"
        "本地运行环境依赖见 GitHub requirements.txt",
        "复用主模型，无额外显存",
        STAGE,
        PAPER,
        REPO,
        OWNER,
    ],
    # --- 技术点 3 -------------------------------------------------------
    [
        "机械 OCR 仪器通道",
        "采用 OCR 文字确定性比对架构，覆盖精确匹配、包含关系、序列相似度、词召回与数字核对，"
        "输出 19 项机械核对特征。",
        '{"claim_text": "string", "ocr_texts": "list[string]", '
        '"ocr_scores": "list[float](optional)"}',
        '{"features": "object(19 项机械特征)", "note": "string"}',
        "使用 GitHub 链接中提供的 qacd.mechanical_ocr_features 接口 + RapidOCR，"
        "本地运行环境依赖见 GitHub requirements.txt",
        "CPU 运行，无显存需求",
        STAGE,
        PAPER,
        REPO,
        OWNER,
    ],
    # --- 技术点 4 -------------------------------------------------------
    [
        "保形选择性预测与弃权",
        "采用 class-conditional conformal p-value + BH/BY 多重性控制架构，"
        "输出覆盖率、正确保留率与 realized FDP，支持按错误代价选择工作点。",
        '{"calibration_null_scores": "list[float]", "test_scores": "list[float]", '
        '"alpha": "float(optional)", "procedure": "string(optional)"}',
        '{"accepted": "list[bool]", "coverage": "float", "realized_fdp": "float|null"}',
        "使用 GitHub 链接中提供的 qacd.conformal_select 接口，"
        "本地运行环境依赖见 GitHub requirements.txt",
        "CPU 运行，无显存需求",
        STAGE,
        PAPER,
        REPO,
        OWNER,
    ],
]

#: The single technology point for 课题2 (基于不确定性量化的风险评估与校准方法).
#: Wording follows the project application SQ2025AAA011084: multi-dimensional
#: uncertainty decoupling, cross-modal semantic fidelity and generation-space
#: uncertainty quantification, risk-uncertainty calibration, and early warning
#: for high-risk generated content.
TOPIC2_ROWS = [
    [
        "基于不确定性量化的多模态生成内容风险评估与校准",
        "采用多维不确定性解耦 + 风险-不确定性映射校准双架构，覆盖跨模态语义保真度、"
        "生成空间不确定性与证据一致性，输出生成内容风险分与校准概率，支持高风险内容提前预警。"
        "黑盒后处理，不改动被评估模型。",
        '{"question": "string", "answer": "string", "image": "string(optional)", '
        '"threshold": "float(optional)"}',
        '{"risk_score": "float(0-1)", "is_high_risk": "bool", '
        '"calibrated_confidence": "float", "num_claims": "int", "version": "string"}',
        "使用 GitHub 链接中提供的 qacd 包与 run.py，本地运行环境依赖见 GitHub requirements.txt",
        "单卡 40G 显存",
        STAGE,
        PAPER,
        REPO,
        "湖南大学（待补充）",
    ],
]

FONT = "等线"
CELL_SIZE = Pt(8)


def set_run_font(run, name: str = FONT, size: Pt = CELL_SIZE, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = size
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def write_cell(cell, text: str, size: Pt = CELL_SIZE, bold: bool = False) -> None:
    """Replace a cell's content with a single formatted paragraph."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold)


def add_heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=Pt(11 if level == 1 else 10), bold=True)


def add_body(document, text: str, size: Pt = Pt(9.5)) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)


def add_bullets(document, items) -> None:
    for item in items:
        # The project template does not define the built-in "List Bullet" style,
        # so the bullet is part of the text rather than a numbering definition.
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run("• " + item)
        set_run_font(run, size=Pt(9.5))


def set_table_borders(table) -> None:
    """Draw single-line borders on every edge.

    ``Table Grid`` is a built-in Word style that the project template's
    ``styles.xml`` does not carry, so the borders are written directly.
    """
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(existing)
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.makeelement(qn(f"w:{edge}"), {})
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def add_simple_table(document, header, rows, widths=None):
    table = document.add_table(rows=1, cols=len(header))
    set_table_borders(table)
    for i, text in enumerate(header):
        write_cell(table.rows[0].cells[i], text, size=Pt(9), bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            write_cell(cells[i], text, size=Pt(9))
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = width
    return table


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--figure", default=None, help="optional main-result figure PNG")
    parser.add_argument(
        "--keep-example-row",
        action="store_true",
        help="keep the worked-example row from the template (row 1) in the table",
    )
    parser.add_argument(
        "--table-only",
        action="store_true",
        help="fill the table and change nothing else in the template",
    )
    parser.add_argument(
        "--topic2",
        action="store_true",
        help="fill the single 课题2 row instead of the four技术点 rows",
    )
    args = parser.parse_args()

    document = docx.Document(args.template)

    # --- locate the 技术集成 table -------------------------------------
    table = document.tables[0]
    if len(table.rows) < 6 or len(table.columns) != 12:
        raise SystemExit(
            f"unexpected template shape: {len(table.rows)} rows x {len(table.columns)} cols"
        )

    # Row 0 is the header, row 1 the template example, rows 2-5 the four blanks.
    # Column 0 already holds the sequence numbers 1..4; column 1 is 所属子课题.
    rows = TOPIC2_ROWS if args.topic2 else ROWS
    for offset, values in enumerate(rows):
        row = table.rows[2 + offset]
        for col, value in enumerate([SUBTopic] + list(values), start=1):
            write_cell(row.cells[col], value)
        # Emphasise the technology name (column 2) for scanability.
        name_run = row.cells[2].paragraphs[0].runs[0]
        name_run.bold = True

    # --table-only: the caller wants the template's own content left alone.
    # Fill the four blank rows, save, and change nothing else.
    if args.table_only:
        document.save(args.out)
        print(f"[docx] table filled -> {args.out}")
        return 0

    # The template ships a worked example in row 1. It is useful while filling
    # the table in and noise in a submission, so drop it unless explicitly kept.
    if not args.keep_example_row:
        table._tbl.remove(table.rows[1]._tr)

    # Repeat the header row when the table spans pages.
    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    tbl_header = tr_pr.makeelement(qn("w:tblHeader"), {})
    tr_pr.append(tbl_header)

    # --- relabel the template marker -----------------------------------
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "【技术集成 模版】":
            for run in paragraph.runs:
                run.text = ""
            run = paragraph.add_run("【技术集成 · 课题一 QACD 填写稿】")
            set_run_font(run, size=Pt(12), bold=True)

    # ==================================================================
    # Supplementary section
    # ==================================================================
    document.add_page_break()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("课题一 QACD 技术对接补充说明")
    set_run_font(run, size=Pt(15), bold=True)

    add_body(
        document,
        "本补充说明随上表一并提交，用于平台建设排期与工作量预估。"
        "所有数值均来自远程研究服务器上的冻结实验产出，未做二次拟合或人工调整。",
    )

    # --- 1 交付状态总览 -------------------------------------------------
    add_heading(document, "一、交付状态总览")
    add_simple_table(
        document,
        ["序号", "技术点", "是否可纳入平台建设", "说明"],
        [
            ["1", "QACD 核心打分器", "可以", "接口已定义；需接入平台 LVLM 与带标签开发集后产出打分器"],
            ["2", "MVR 多视图重采样通道", "可以", "依赖同一 LVLM 的多次采样能力"],
            ["3", "机械 OCR 仪器通道", "可以", "无 GPU 依赖，可独立验证"],
            ["4", "保形选择性预测与弃权", "可以", "纯 CPU 计算，工作点在平台自有校准批次上标定"],
        ],
        widths=[Inches(0.5), Inches(1.7), Inches(1.3), Inches(2.8)],
    )

    # --- 2 关键数值 -----------------------------------------------------
    add_heading(document, "二、关键结果（冻结 TextVQA 测试集）")
    add_body(
        document,
        "测试集规模：1,999 条回答 / 1,278 张开发未见图像 / 836 条失败（41.8%）。"
        "配对图像级 bootstrap 5,000 次，随机种子固定。",
    )
    add_simple_table(
        document,
        ["方法", "回答级 AUROC", "图像级 AUROC", "备注"],
        [
            ["UMPIRE K=5", "0.8564", "0.8561", "白盒，需要模型内部量"],
            ["QACD（LM + 机械仪器 + MVR K=5）", "0.8526", "—", "黑盒，本课题方法"],
            ["SelfCheckGPT-NLI K=5", "0.8319", "0.8477", "公开基线"],
            ["Multi-sample Consistency K=5", "0.8258", "0.8363", "公开基线"],
            ["QACD（原始 95 维配置）", "0.7991", "—", "早期配置"],
        ],
        widths=[Inches(2.5), Inches(1.1), Inches(1.1), Inches(1.6)],
    )
    add_body(
        document,
        "配对对比：vs UMPIRE K=5 为 −0.0039（95% CI [−0.0195, +0.0124]，区间跨 0，统计持平）；"
        "vs SelfCheckGPT-NLI 为 +0.0207（[+0.0028, +0.0391]）；"
        "vs Multi-sample Consistency 为 +0.0267（[+0.0102, +0.0436]）。"
        "即：黑盒 QACD 与需要模型内部量的白盒方法统计持平，并优于两个公开采样类基线。",
    )
    add_body(
        document,
        "上述数值可复现：仓库执行 python scripts/run_frozen_evaluation.py 即用其自身的校准器、"
        "聚合与指标在冻结证据上重算全部结果，并与冻结重放记录逐臂比对——主张级两臂绝对差 1e-6，"
        "三个公开基线完全吻合。比对明细见仓库 results/reproduction_check.csv，"
        "证据包与输出的 SHA256 记录在 results/manifest.json。",
    )
    if args.figure and Path(args.figure).exists():
        document.add_picture(args.figure, width=Inches(6.2))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(
            "图 1  左：冻结测试集上的回答级 AUROC 对比；右：配对图像级 bootstrap 的 ΔAUROC 与 95% 置信区间"
        )
        set_run_font(run, size=Pt(8.5))

    # --- 3 资源与工作量 -------------------------------------------------
    add_heading(document, "三、资源需求与工作量预估")
    add_bullets(
        document,
        [
            "显存：被评估模型 LLaVA-1.5-13B 在 fp16 下约 26 GB，建议单卡 ≥40 GB（A100-40G/80G 等）。"
            "24 GB 卡可考虑 8-bit/4-bit 量化或改用 7B 主干。",
            "决策层：仅依赖 NumPy，CPU 可运行，无 GPU 需求。",
            "模型调用次数（算力成本代理）：LM 通道本身 7.65 次/回答；"
            "4 视图 + MVR K=5 为 12.61 次；每增加 1 层采样 +1 次。",
            "成本优化建议：视图数由 4 降至 1，精度无统计可辨差异（AUROC 0.8509 → 0.8515，区间跨 0），"
            "而调用次数由 12.61 降至 9.61（K=5 时减少 24%）。受算力约束时优先降视图数，不要降 K。",
            "开发集规模：学习曲线显示 300 条尚未饱和，增益大部分在约 1,200 条时已实现。"
            "平台接入新领域建议准备 ≥1,200 条带标签回答再冻结打分器。",
        ],
    )

    # --- 4 平台必须拦截 -------------------------------------------------
    add_heading(document, "四、平台侧必须拦截的两种情况")
    add_simple_table(
        document,
        ["情况", "判别方式", "平台动作"],
        [
            ["打分器未拟合", "GET /health 返回 scorer_fitted:false，或响应 warnings 含 "
                              "\"scorer is not fitted\"", "拒绝将 risk_score 写入业务库"],
            ["特征维度不匹配", "响应 feature_dim 与平台记录不一致", "拒绝并告警（表示打分器版本变更）"],
        ],
        widths=[Inches(1.3), Inches(3.0), Inches(2.0)],
    )

    # --- 5 集成顺序 -----------------------------------------------------
    add_heading(document, "五、建议的集成顺序")
    add_bullets(
        document,
        [
            "第一步：接技术点 3（机械 OCR 通道）——无 GPU 依赖，可独立验证，最容易打通链路。",
            "第二步：接技术点 1（核心打分器）——需要 LVLM 运行环境与平台自有开发集。",
            "第三步：加技术点 2（MVR 通道）作为精度增量，K 从 3 起调。",
            "第四步：接技术点 4（保形弃权），在平台自有校准批次上标定工作点。",
        ],
    )

    # --- 6 待补充事项 ---------------------------------------------------
    add_heading(document, "六、待补充事项")
    add_bullets(
        document,
        [
            "论文尚未投稿，暂无 arXiv 链接；待投稿后补充。",
            "作者信息、单位与通信作者信息待补充。",
            "对接负责人姓名待确认。",
        ],
    )

    # --- 附：仓库指引 ---------------------------------------------------
    add_heading(document, "附：代码与文档入口")
    add_body(
        document,
        f"代码仓库：{REPO}　（含中英双语文档、可运行的参考实现、128 项离线测试与接口文档）",
    )
    add_bullets(
        document,
        [
            "docs/01-方法说明.md —— 方法、流程与结果",
            "docs/02-接口文档.md —— 四个对接端点的入参/出参契约",
            "docs/03-部署与资源需求.md —— 依赖、显存、调用次数、部署步骤",
            "docs/04-平台对接说明.md —— 对接边界、验收标准、监控与常见问题",
            "docs/05-课题一技术对接表.md —— 本表的 Markdown 镜像",
            "results/ —— 冻结实验数值与溯源",
        ],
    )

    document.save(args.out)
    print(f"[docx] written -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
